from pptx import Presentation
from pptx.util import Inches

commands = [
("configure", "Configure goose settings"),
("info", "Display goose information"),
("doctor", "Check that your Goose setup is working"),
("mcp", "Run one of the mcp servers bundled with goose"),
("acp", "Run goose as an ACP agent server on stdio"),
("serve", "Start ACP server over HTTP and WebSocket"),
("session (alias: s)", "Start or resume interactive chat sessions"),
("project (alias: p)", "Open the last project directory"),
("projects (alias: ps)", "List recent project directories"),
("run", "Execute commands from an instruction file or stdin"),
("recipe", "Recipe utilities for validation and deeplinking"),
("skills", "Skill utilities"),
("plugin", "Manage plugins"),
("schedule (alias: sched)", "Manage scheduled jobs"),
("gateway (alias: gw)", "Manage gateways for external platform integrations"),
("update", "Update the goose CLI version"),
("term", "Terminal-integrated goose session"),
("tui", "Launch the goose terminal UI"),
("local-models (alias: lm)", "Manage local inference models"),
("completion", "Generate the autocompletion script or Nushell module for the specified shell"),
("review", "Review the current diff using goose"),
("help", "Print this message or help for a given subcommand"),
]

global_opts = [
"-h, --help — Print help",
"-V, --version — Print version",
]

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "goose CLI — Supported Commands"

left, top, width, height = Inches(0.5), Inches(1.4), Inches(12.0), Inches(5.2)
box = slide.shapes.add_textbox(left, top, width, height)

tf = box.text_frame

tf.word_wrap = True

tf.clear()

p = tf.paragraphs[0]
p.text = "Commands:"

for name, desc in commands:
    para = tf.add_paragraph()
    para.text = f"• {name}: {desc}"
    para.level = 1

para = tf.add_paragraph()
para.text = "\nGlobal options:"

for opt in global_opts:
    po = tf.add_paragraph()
    po.text = f"• {opt}"
    po.level = 1

out = "$OUTDIR/goose_help_slide.pptx"
import os
out = os.path.join(os.path.dirname(__file__), "goose_help_slide.pptx")
prs.save(out)
print(out)
